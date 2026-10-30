"""The settings you edit, and the reconciliation report the generator writes.

Two kinds of settings, both optional:

  1. Global exclusions: settings/exclusions.yaml
     Column names that are never compared, for every table. One YAML file. Example: the Fivetran metadata columns.

  2. Table rules: the Excel workbook in settings/ (the one file there is used every run)
     One row per table. Columns:
       table_name     the source table name (required)
       key_columns    the key, comma separated, when the table has no primary key or it should not be used
       exclude_columns  columns to leave out for this table only, comma separated
       source_where   a row filter for the SOURCE query, written as SQL without the word WHERE
       target_where   a row filter for the SNOWFLAKE query, written as SQL without the word WHERE. {env} is allowed in it.
       incremental_column  the column an incremental run windows on (default updated_at). Optional column.
     The same filter is applied to the count and to the data query. A join is written as a condition with a
     sub-select, for example:  customer_id IN (SELECT id FROM public.customers WHERE active = true)

Without a rules file, nothing is filtered and no table has a key override. That is the default, and it is a
normal run, not an error.
"""

import csv
import sys
from dataclasses import dataclass, field
from pathlib import Path

import yaml

from yaml_generation.yaml_builder import check_filter, normalize_name

SETTINGS_DIR = Path(__file__).resolve().parent / "settings"
GLOBAL_EXCLUSIONS_FILE = SETTINGS_DIR / "exclusions.yaml"
RULES_SHEET = "table_rules"
RULE_HEADERS = ["table_name", "key_columns", "exclude_columns", "source_where", "target_where", "incremental_column"]
# The column an incremental run windows on, when the rules workbook does not name one. The validators' freshness check
# reads this column too, so a table without it cannot run an incremental window.
INCREMENTAL_COLUMN = "updated_at"


@dataclass
class TableRule:
    """What the rules workbook says about one table. Empty values mean 'no rule'."""

    key_columns: list = None
    exclude: set = field(default_factory=set)
    source_where: str = None
    target_where: str = None
    incremental_column: str = None


def _names(text) -> list:
    """'Company_ID, site id' -> ['company_id', 'site_id']. Names are normalised like the columns are."""
    if not text:
        return []
    return [normalize_name(part) for part in str(text).split(",") if part.strip()]


def load_global_exclusions(path=GLOBAL_EXCLUSIONS_FILE) -> set:
    """Normalised names of columns excluded for every table. A missing file means no global exclusions."""
    if not Path(path).exists():
        return set()
    data = yaml.safe_load(Path(path).read_text(encoding="utf-8")) or {}
    return {normalize_name(str(name)) for name in (data.get("exclude_columns") or [])}


def load_table_rules(path, source_type=None) -> dict:
    """Read the rules workbook into {table_name: TableRule}. Raises ValueError with the row number if something is wrong.

    source_type is the --source_type of the run. It is needed to check the source_where column.
    """
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook[RULES_SHEET] if RULES_SHEET in workbook.sheetnames else workbook.active
        rows = list(sheet.iter_rows(values_only=True))
    finally:
        workbook.close()
    if not rows:
        return {}

    headers = [str(h or "").strip().lower() for h in rows[0]]
    unknown = [h for h in headers if h and h not in RULE_HEADERS]
    if unknown or "table_name" not in headers:
        raise ValueError(f"{Path(path).name}: the first row must hold the headers {RULE_HEADERS}; "
                         f"found {[h for h in headers if h]}" + (f", unknown: {unknown}" if unknown else ""))
    position = {name: headers.index(name) for name in RULE_HEADERS if name in headers}

    rules = {}
    for number, row in enumerate(rows[1:], start=2):
        def cell(name):
            index = position.get(name)
            value = row[index] if index is not None and index < len(row) else None
            return None if value is None or not str(value).strip() else str(value).strip()

        table = cell("table_name")
        if table is None:
            continue
        if table in rules:
            raise ValueError(f"{Path(path).name}, row {number}: table '{table}' is listed twice")
        rules[table] = TableRule(
            key_columns=_names(cell("key_columns")) or None,
            exclude=set(_names(cell("exclude_columns"))),
            source_where=check_filter(cell("source_where"), f"row {number} ({table}) source_where", source_type),
            target_where=check_filter(cell("target_where"), f"row {number} ({table}) target_where", placeholders=True),
            incremental_column=cell("incremental_column"),
        )
    return rules


def pick_rules_file(explicit=None, folder=SETTINGS_DIR, ask=input, interactive=None):
    """The rules workbook to use, or None.

    --rules_file wins. Otherwise the .xlsx files in settings/ are offered: one file is used directly, several
    give a numbered list to choose from, and none means no rules. Without a terminal to answer, several files
    are an error, so a scheduled run never picks one by accident.
    """
    if explicit:
        path = Path(explicit)
        if not path.exists():
            raise FileNotFoundError(f"rules file not found: {path}")
        return path
    files = sorted(p for p in Path(folder).glob("*.xlsx") if not p.name.startswith("~$"))
    if not files:
        return None
    if len(files) == 1:
        return files[0]
    if interactive is None:
        interactive = sys.stdin.isatty()
    if not interactive:
        raise ValueError(f"several rules files in {folder}; keep only the table rules workbook there")
    print(f"Rules files in {folder}:")
    for number, file in enumerate(files, start=1):
        print(f"  {number}. {file.name}")
    answer = ask("Number of the file to use (press Enter for none): ").strip()
    if not answer:
        return None
    if not answer.isdigit() or not 1 <= int(answer) <= len(files):
        raise ValueError(f"'{answer}' is not a number from the list")
    return files[int(answer) - 1]


RECONCILIATION_HEADERS = ["run_at", "table", "source_column", "target_column", "status", "note"]


def write_reconciliation(path, run_at: str, rows) -> Path:
    """Write the reconciliation report: one line per column (or per table), with its status.

    rows = [(table, source_column, target_column, status, note), ...]
    """
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    with open(path, "w", encoding="utf-8", newline="") as f:
        writer = csv.writer(f)
        writer.writerow(RECONCILIATION_HEADERS)
        for table, source_column, target_column, status, note in rows:
            writer.writerow([run_at, table, source_column, target_column, status, note])
    return path
