"""Integration tests for generate_yaml.py: the whole generation for one run, against fake source and Snowflake databases.

Each test writes into a temporary folder, so the real config/ and output/ folders are never touched. The fakes are
described in helpers.py. The live PostgreSQL run is in test_live_postgres.py.
"""

import logging
from types import SimpleNamespace

import pytest
import yaml
from openpyxl import Workbook

from helpers import (LOG, PRIMARY_KEYS, SOURCE_SCHEMA, SOURCE_TABLES, TARGET_SCHEMA, TARGET_TABLES, FakeDB, TableRule,
                     config_dir_for, make_args, run_generate, statuses)
from yaml_generation import filter_files, generate_yaml, settings
from yaml_generation import yaml_builder as yb


def read_yaml(path):
    return yaml.safe_load(path.read_text(encoding="utf-8"))


def data_block(config, table):
    return read_yaml(config / "data_validation" / f"{table}.yaml")["tables"][table]["validations"]["data_validation"]


def count_block(config, layer_file, table):
    return read_yaml(config / "count_validation" / f"{layer_file}.yaml")["tables"][table]["validations"]["count_validation"]


def detail_of(results, table, kind):
    return next(detail for t, k, _, detail in results if t == table and k == kind)


# --- resolving the arguments -----------------------------------------------------------------------------------

def test_defaults_come_from_the_rules_and_the_environment():
    args = generate_yaml.resolve_args(SimpleNamespace(source_type="postgres", tables=["units"], environment="dev",
                                                      data_validation="yes", count_validation="yes", source_schema=None,
                                                      target_database=None, target_schema=None))
    assert (args.source_schema, args.target_schema, args.target_database) == \
        ("public", "STOREDGE_FMS_PUBLIC", "DEV_EDGE_BRONZE")
    assert args.target_template == "{env}_EDGE_BRONZE"


def test_local_reads_the_dev_database_and_writes_the_placeholder():
    args = make_args(environment="local")
    assert (args.target_database, args.target_template) == ("DEV_EDGE_BRONZE", "{env}_EDGE_BRONZE")


def test_an_unknown_environment_is_refused():
    with pytest.raises(ValueError, match="unknown --environment uat"):
        make_args(environment="uat")


def test_a_target_database_equal_to_the_environment_name_is_accepted():
    assert make_args(target_database="DEV_EDGE_BRONZE").target_database == "DEV_EDGE_BRONZE"


def test_a_target_database_of_another_environment_is_refused():
    with pytest.raises(ValueError, match="does not match --environment"):
        make_args(target_database="PRD_EDGE_BRONZE")


def test_athena_needs_a_target_database_and_writes_it_as_given():
    """Athena and Redshift have no Snowflake name of their own, so --target_database gives it, unchanged."""
    with pytest.raises(ValueError, match="--target_database is required for Athena"):
        make_args(source_type="athena", source_schema="glue_db")
    args = make_args(source_type="athena", source_schema="glue_db", target_database="ANY_DB")
    assert (args.target_database, args.target_template) == ("ANY_DB", "ANY_DB")


def test_all_cannot_be_combined_with_table_names():
    with pytest.raises(ValueError, match="'all' cannot be combined"):
        make_args(tables=("all", "units"))


def test_nothing_to_generate_is_refused():
    with pytest.raises(ValueError, match="nothing to generate"):
        make_args(data="no", count="no")


@pytest.mark.xfail(strict=True, reason="ISSUE-02: athena with no --source_schema is accepted; the run then reports every table as missing from schema None")
def test_athena_without_a_schema_is_refused_before_any_table_is_read():
    args = SimpleNamespace(source_type="athena", tables=["all"], environment="dev", data_validation="yes",
                           count_validation="yes", source_schema=None, target_database="DEV_X_BRONZE",
                           target_schema=None, run_type="historical")
    with pytest.raises(ValueError):
        generate_yaml.resolve_args(args)


# --- a historical run ------------------------------------------------------------------------------------------

def test_historical_run_writes_the_expected_kinds_and_statuses(tmp_path):
    config = config_dir_for(tmp_path, "historical")
    results, _ = run_generate(make_args("historical"), config)
    assert statuses(results) == {
        ("lookup", "data_validation"): "FAILED",
        ("notes", "data_validation"): "NEW", ("notes", "count_validation"): "NEW",
        ("units", "data_validation"): "NEW", ("units", "count_validation"): "NEW",
    }
    assert detail_of(results, "lookup", "data_validation").startswith(
        "unable to find the primary key of edge.lookup in the source")


def test_a_table_with_no_primary_key_fails_and_the_log_says_why(tmp_path, caplog):
    with caplog.at_level(logging.ERROR):
        results, _ = run_generate(make_args("historical", tables=("lookup",)), config_dir_for(tmp_path, "historical"))
    assert statuses(results) == {("lookup", "data_validation"): "FAILED"}, "nothing is written for a failed table, its count too"
    assert "unable to find the primary key of edge.lookup" in caplog.text


def test_the_yaml_holds_the_env_placeholder_on_the_target_side_only(tmp_path):
    config = config_dir_for(tmp_path, "historical")
    run_generate(make_args("historical"), config)
    block = data_block(config, "units")
    count = count_block(config, "bronze_postgres", "units")
    assert block["target_database"] == count["target_database"] == "{env}_EDGE_BRONZE"
    assert "FROM {env}_EDGE_BRONZE." in block["targetquery"] and "FROM {env}_EDGE_BRONZE." in count["targetquery"]
    assert "{" not in block["sourcequery"] and "{" not in count["sourcequery"]
    for text in ((config / "data_validation" / "units.yaml").read_text(encoding="utf-8"),
                 (config / "count_validation" / "bronze_postgres.yaml").read_text(encoding="utf-8")):
        assert "DEV_EDGE_BRONZE" not in text, "the environment is filled in by the validators, not written here"


def test_historical_data_file_drops_the_global_exclusions_and_lists_them_in_the_header(tmp_path):
    config = config_dir_for(tmp_path, "historical")
    run_generate(make_args("historical"), config)
    assert "_fivetran" not in data_block(config, "units")["sourcequery"].lower()
    header = (config / "data_validation" / "units.yaml").read_text(encoding="utf-8")
    assert "# Key: id_normalized   Columns compared: 4" in header
    assert "# Excluded (global list): _fivetran_synced, _fivetran_active" in header


def test_second_historical_run_leaves_every_file_unchanged(tmp_path):
    config = config_dir_for(tmp_path, "historical")
    run_generate(make_args("historical"), config)
    results, _ = run_generate(make_args("historical"), config)
    got = statuses(results)
    assert got[("units", "data_validation")] == "UNCHANGED"
    assert got[("notes", "data_validation")] == "UNCHANGED"
    assert got[("units", "count_validation")] == "UNCHANGED"


def test_a_source_column_with_no_target_match_is_compared_as_missing_not_dropped(tmp_path):
    """A column with no match in Snowflake stays in the comparison. The Snowflake side holds a token, so every row differs."""
    config = config_dir_for(tmp_path, "historical")
    run_generate(make_args("historical"), config)
    changed = {**SOURCE_TABLES, ("edge", "units"): SOURCE_TABLES[("edge", "units")] + [("note_text", "text")]}
    results, _ = run_generate(make_args("historical"), config, source=FakeDB(changed, PRIMARY_KEYS))
    got = statuses(results)
    assert got[("units", "data_validation")] == "UPDATED"
    assert got[("units", "count_validation")] == "UNCHANGED"
    block = data_block(config, "units")
    assert "note_text_normalized" in block["sourcequery"]
    assert f"CAST('{yb.MISSING_IN_TARGET}' AS STRING) AS note_text_normalized" in block["targetquery"]
    assert (f"# Source columns not in Snowflake (compared as {yb.MISSING_IN_TARGET}, so every row differs): note_text"
            in (config / "data_validation" / "units.yaml").read_text(encoding="utf-8"))


def test_a_snowflake_column_with_no_source_match_is_compared_as_missing_not_dropped(tmp_path):
    config = config_dir_for(tmp_path, "historical")
    run_generate(make_args("historical"), config)
    target = {**TARGET_TABLES, ("EDGE", "UNITS"): TARGET_TABLES[("EDGE", "UNITS")] + [("EXTRA_FLAG", "BOOLEAN")]}
    results, _ = run_generate(make_args("historical"), config, target=FakeDB(target))
    assert statuses(results)[("units", "data_validation")] == "UPDATED"
    block = data_block(config, "units")
    assert f"CAST('{yb.MISSING_IN_SOURCE}' AS TEXT) AS extra_flag_normalized" in block["sourcequery"]
    assert '"EXTRA_FLAG"' in block["targetquery"]
    assert (f"# Snowflake columns not in source (compared as {yb.MISSING_IN_SOURCE}, so every row differs): EXTRA_FLAG"
            in (config / "data_validation" / "units.yaml").read_text(encoding="utf-8"))


def test_the_global_list_is_what_drops_a_column(tmp_path):
    config = config_dir_for(tmp_path, "historical")
    run_generate(make_args("historical"), config, global_excluded=set())
    assert "_fivetran_synced_normalized" in data_block(config, "units")["sourcequery"]


def test_reconciliation_has_a_status_for_every_source_and_target_column(tmp_path):
    _, rows = run_generate(make_args("historical"), config_dir_for(tmp_path, "historical"))
    units = {(s, t): status for table, s, t, status, _ in rows if table == "units"}
    assert units[("_fivetran_synced", "_FIVETRAN_SYNCED")] == "EXCLUDED_GLOBAL"
    assert units[("id", "ID")] == "MATCHED"


def test_a_table_rule_excludes_a_column_and_filters_both_sides_for_that_table_only(tmp_path):
    config = config_dir_for(tmp_path, "historical")
    rules = {"units": TableRule(exclude={"unit_name"}, source_where="id > 0", target_where="ID > 0")}
    run_generate(make_args("historical"), config, table_rules=rules)
    block = data_block(config, "units")
    assert "unit_name" not in block["sourcequery"].lower()
    assert "WHERE (id > 0)" in block["sourcequery"] and "(ID > 0)" in block["targetquery"]
    assert count_block(config, "bronze_postgres", "units")["sourcequery"].endswith("WHERE (id > 0);")
    assert "# Excluded (table rule): unit_name" in (config / "data_validation" / "units.yaml").read_text(encoding="utf-8")
    assert "WHERE (" not in data_block(config, "notes")["sourcequery"], "the rule is for one table only"


def test_an_excluded_key_column_fails_the_data_file_and_writes_no_count(tmp_path):
    results, _ = run_generate(make_args("historical"), config_dir_for(tmp_path, "historical"),
                              table_rules={"notes": TableRule(exclude={"id"})})
    assert statuses(results)[("notes", "data_validation")] == "FAILED"
    assert detail_of(results, "notes", "data_validation").startswith("key column(s) not in both source and Snowflake: id")
    assert ("notes", "count_validation") not in statuses(results)


def test_a_key_missing_from_snowflake_fails_the_data_file(tmp_path):
    columns = [c for c in TARGET_TABLES[("EDGE", "UNITS")] if c[0] != "ID"]
    results, _ = run_generate(make_args("historical", tables=("units",)), config_dir_for(tmp_path, "historical"),
                              target=FakeDB({**TARGET_TABLES, ("EDGE", "UNITS"): columns}))
    assert statuses(results)[("units", "data_validation")] == "FAILED"
    assert detail_of(results, "units", "data_validation").startswith("key column(s) not in both source and Snowflake: id")
    assert ("units", "count_validation") not in statuses(results)


def test_a_key_rule_gives_a_key_to_a_table_without_one(tmp_path):
    config = config_dir_for(tmp_path, "historical")
    results, _ = run_generate(make_args("historical"), config, table_rules={"lookup": TableRule(key_columns=["code"])})
    assert statuses(results)[("lookup", "data_validation")] == "NEW"
    assert data_block(config, "lookup")["sourcecolumn"] == "code_normalized"


def test_one_table_failing_does_not_stop_the_others(tmp_path):
    config = config_dir_for(tmp_path, "historical")
    source = FakeDB(SOURCE_TABLES, PRIMARY_KEYS, fail_on=("notes",))
    results, _ = run_generate(make_args("historical"), config, source=source)
    got = statuses(results)
    assert got[("notes", "table")] == "FAILED"
    assert got[("units", "data_validation")] == "NEW"
    assert set(yb._count_tables(config, "bronze_postgres")) == {"units"}, "notes failed and lookup has no key: neither is counted"


def test_a_table_missing_from_the_source_is_an_error_and_its_yaml_is_kept(tmp_path):
    config = config_dir_for(tmp_path, "historical")
    run_generate(make_args("historical", tables=("notes",)), config)
    gone = FakeDB({k: v for k, v in SOURCE_TABLES.items() if k[1] != "notes"}, PRIMARY_KEYS)
    results, _ = run_generate(make_args("historical", tables=("notes",)), config, source=gone)
    assert statuses(results)[("notes", "table")] == "FAILED"
    assert "nothing was changed" in detail_of(results, "notes", "table")
    assert (config / "data_validation" / "notes.yaml").exists(), "nothing is deleted"
    assert "notes" in yb._count_tables(config, "bronze_postgres")


def test_a_table_missing_from_snowflake_is_an_error_and_nothing_is_deleted(tmp_path):
    config = config_dir_for(tmp_path, "historical")
    run_generate(make_args("historical"), config)
    target = FakeDB({k: v for k, v in TARGET_TABLES.items() if k[1] != "NOTES"})
    results, _ = run_generate(make_args("historical"), config, target=target)
    assert statuses(results)[("notes", "table")] == "FAILED"
    assert (config / "data_validation" / "notes.yaml").exists()
    assert "notes" in yb._count_tables(config, "bronze_postgres")


def test_all_leaves_a_generated_table_that_left_the_source_alone(tmp_path, caplog):
    config = config_dir_for(tmp_path, "historical")
    run_generate(make_args("historical"), config)
    gone = FakeDB({k: v for k, v in SOURCE_TABLES.items() if k[1] != "notes"}, PRIMARY_KEYS)
    results, _ = run_generate(make_args("historical"), config, source=gone)
    assert not any(status == "REMOVED" for _, _, status, _ in results)
    assert (config / "data_validation" / "notes.yaml").exists()
    assert "notes" in yb._count_tables(config, "bronze_postgres")


def test_all_never_deletes_a_file_the_generator_did_not_write(tmp_path):
    config = config_dir_for(tmp_path, "historical")
    manual = config / "data_validation" / "manual.yaml"
    manual.parent.mkdir(parents=True)
    manual.write_text("tables: {}\n", encoding="utf-8")
    run_generate(make_args("historical"), config)
    assert manual.exists()


# --- an incremental run writes to its own folder -----------------------------------------------------------------

def test_an_incremental_run_does_not_touch_the_historical_folder(tmp_path):
    run_generate(make_args("historical"), config_dir_for(tmp_path, "historical"))
    before = (config_dir_for(tmp_path, "historical") / "data_validation" / "units.yaml").read_text(encoding="utf-8")
    run_generate(make_args("incremental", tables=("all",)), config_dir_for(tmp_path, "incremental"))
    assert (config_dir_for(tmp_path, "historical") / "data_validation" / "units.yaml").read_text(encoding="utf-8") == before


# --- row filters from the workbooks ------------------------------------------------------------------------------

def test_a_rules_target_where_with_the_env_placeholder_is_written_as_it_is(tmp_path):
    """The rules workbook accepts {env} in target_where. The target query carries it as written."""
    path = tmp_path / "rules.xlsx"
    book = Workbook()
    ws = book.active
    ws.title = settings.RULES_SHEET
    ws.append(settings.RULE_HEADERS)
    ws.append(["units", None, None, None, "ID IN (SELECT ID FROM {env}_EDGE_BRONZE.EDGE.NOTES)"])
    book.save(path)
    rules = settings.load_table_rules(path, "postgres")
    config = config_dir_for(tmp_path, "historical")
    results, _ = run_generate(make_args("historical", tables=("units",)), config, table_rules=rules)
    assert statuses(results)[("units", "data_validation")] == "NEW"
    assert "(ID IN (SELECT ID FROM {env}_EDGE_BRONZE.EDGE.NOTES))" in data_block(config, "units")["targetquery"]


def test_a_file_filter_with_the_env_placeholder_on_the_target_side_is_written_and_checked(tmp_path):
    """The filter workbook keeps {env} on the Snowflake side. The target query holds it, and _checked() accepts it."""
    book = Workbook()
    ws = book.active
    ws.title = "filters"
    ws.append(["Target Table Name", "Historical Filter to apply."])
    ws.append(["units", "units j\njoin notes n on j.note_id = n.id\nwhere n.body is not null"])
    book.save(tmp_path / "filters.xlsx")
    file_filters = filter_files.load("postgres", SOURCE_SCHEMA, "{env}_EDGE_BRONZE", "EDGE", "historical",
                                     folder=tmp_path, log=LOG)
    config = config_dir_for(tmp_path, "historical")
    results, _ = run_generate(make_args("historical", tables=("units",)), config, file_filters=file_filters)
    assert statuses(results)[("units", "data_validation")] == "NEW", "a target query that fails _checked() would be FAILED"
    assert ('note_id IN (SELECT id FROM {env}_EDGE_BRONZE."EDGE"."NOTES" WHERE body is not null AND _FIVETRAN_ACTIVE = TRUE)'
            in data_block(config, "units")["targetquery"])


# --- run_for_validation: the entry point main.py calls ---------------------------------------------------------

@pytest.fixture
def run_env(monkeypatch, tmp_path):
    """Points the entry point at a temporary folder and fake databases. Returns the folder and a rules workbook."""
    monkeypatch.setattr(generate_yaml, "FRAMEWORK_ROOT", tmp_path)
    monkeypatch.setattr(generate_yaml, "OUTPUT_DIR", tmp_path / "output")
    dbs = {"source": FakeDB(SOURCE_TABLES, PRIMARY_KEYS), "target": FakeDB(TARGET_TABLES)}
    monkeypatch.setattr(generate_yaml, "get_database",
                        lambda db_type, root, env: dbs["source"] if db_type == "postgres" else dbs["target"])
    monkeypatch.setattr(filter_files, "load", lambda *a, **k: {})
    rules = tmp_path / "rules.xlsx"
    book = Workbook()
    ws = book.active
    ws.title = settings.RULES_SHEET
    ws.append(settings.RULE_HEADERS)
    ws.append(["units", None, None, None, "ID > 5"])
    book.save(rules)
    return SimpleNamespace(root=tmp_path, dbs=dbs, rules=rules)


def _run(env, **overrides):
    options = dict(layer_type="bronze_postgres", tables=["all"], environment="dev", run_type="historical",
                   data_validation="yes", count_validation="yes", source_schema=SOURCE_SCHEMA, target_schema=TARGET_SCHEMA,
                   rules_file=env.rules, use_filter_files=False, log=LOG)
    options.update(overrides)
    return generate_yaml.run_for_validation(**options)


def test_run_for_validation_returns_results_and_the_tables_ready_to_validate(run_env):
    results, ready = _run(run_env)
    assert statuses(results)[("units", "data_validation")] == "NEW"
    assert set(ready) == {"notes", "units"}, "lookup has no primary key, so it FAILED and is not ready"
    assert list((run_env.root / "output").glob("reconciliation_*.csv"))


def test_run_for_validation_writes_the_table_rule_from_the_workbook(run_env):
    _run(run_env)
    block = data_block(config_dir_for(run_env.root, "historical"), "units")
    assert "(ID > 5)" in block["targetquery"]


def test_a_table_that_failed_is_not_ready(run_env, monkeypatch):
    failing = FakeDB(SOURCE_TABLES, PRIMARY_KEYS, fail_on=("notes",))
    monkeypatch.setattr(generate_yaml, "get_database",
                        lambda db_type, root, env: failing if db_type == "postgres" else run_env.dbs["target"])
    results, ready = _run(run_env)
    assert statuses(results)[("notes", "table")] == "FAILED"
    assert "notes" not in ready and "units" in ready


def test_a_connection_that_cannot_open_is_a_value_error(run_env, monkeypatch):
    def refuse(db_type, root, env):
        raise FileNotFoundError("creds/dev.yaml")
    monkeypatch.setattr(generate_yaml, "get_database", refuse)
    with pytest.raises(ValueError, match="Cannot open the connections"):
        _run(run_env)


def test_a_bad_rules_workbook_is_a_settings_problem(run_env):
    book = Workbook()
    book.active.title = settings.RULES_SHEET
    book.active.append(["table_name", "bogus"])
    book.save(run_env.rules)
    with pytest.raises(ValueError, match="Settings problem"):
        _run(run_env)


def test_a_file_filter_is_used_only_where_the_rules_workbook_has_none(run_env, monkeypatch):
    monkeypatch.setattr(filter_files, "load", lambda *a, **k: {"units": ("id > 0", "ID > 1")})
    _run(run_env, use_filter_files=True)
    block = data_block(config_dir_for(run_env.root, "historical"), "units")
    assert "(id > 0)" in block["sourcequery"], "no rule for the source side, so the file's filter is used"
    assert "(ID > 5)" in block["targetquery"] and "ID > 1" not in block["targetquery"], "the rules workbook wins"


def test_filter_files_are_not_read_without_the_flag(run_env, monkeypatch):
    def must_not_run(*a, **k):
        raise AssertionError("filter files read without --filter_files yes")
    monkeypatch.setattr(filter_files, "load", must_not_run)
    _run(run_env, use_filter_files=False)


@pytest.mark.xfail(strict=True, reason="ISSUE-03: an unsupported layer raises KeyError, which main.py does not catch as a settings problem")
def test_an_unsupported_layer_is_a_value_error(run_env):
    with pytest.raises(ValueError):
        _run(run_env, layer_type="bronze_silver")


# --- safety: nothing is deleted on a bad listing, or on a flag that does not ask for it ----------------------------

def test_an_empty_source_schema_is_refused_and_deletes_nothing(tmp_path):
    config = config_dir_for(tmp_path, "historical")
    run_generate(make_args("historical"), config)
    empty = FakeDB({}, PRIMARY_KEYS)
    with pytest.raises(ValueError, match="no tables visible in source schema"):
        run_generate(make_args("historical"), config, source=empty)
    assert (config / "data_validation" / "notes.yaml").exists()
    assert "notes" in yb._count_tables(config, "bronze_postgres")


def test_a_run_never_deletes_a_yaml_file_or_a_count_entry(tmp_path):
    config = config_dir_for(tmp_path, "historical")
    run_generate(make_args("historical"), config)
    gone = FakeDB({k: v for k, v in SOURCE_TABLES.items() if k[1] != "notes"}, PRIMARY_KEYS)
    run_generate(make_args("historical", data="no"), config, source=gone)
    assert (config / "data_validation" / "notes.yaml").exists()
    assert "notes" in yb._count_tables(config, "bronze_postgres")


def test_a_count_run_with_count_no_keeps_the_count_entry(tmp_path):
    config = config_dir_for(tmp_path, "historical")
    run_generate(make_args("historical"), config)
    gone = FakeDB({k: v for k, v in SOURCE_TABLES.items() if k[1] != "notes"}, PRIMARY_KEYS)
    run_generate(make_args("historical", count="no", tables=("notes",)), config, source=gone)
    assert "notes" in yb._count_tables(config, "bronze_postgres"), "--count_validation no must not remove the count entry"


def test_an_empty_target_schema_is_refused_and_deletes_nothing(tmp_path):
    # A wrong or mis-cased --target_schema lists no tables in Snowflake. Every generated file would look missing.
    config = config_dir_for(tmp_path, "historical")
    run_generate(make_args("historical"), config)
    with pytest.raises(ValueError, match="no tables visible in Snowflake schema"):
        run_generate(make_args("historical"), config, target=FakeDB({}, PRIMARY_KEYS))
    assert (config / "data_validation" / "notes.yaml").exists()
    assert "notes" in yb._count_tables(config, "bronze_postgres")


def test_a_run_for_another_source_schema_does_not_remove_these_tables(tmp_path):
    config = config_dir_for(tmp_path, "historical")
    run_generate(make_args("historical"), config)
    sales = FakeDB({("sales", "orders"): [("id", "integer")]}, {("sales", "orders"): ["id"]})
    run_generate(make_args("historical", source_schema="sales"), config, source=sales)
    assert (config / "data_validation" / "units.yaml").exists(), "units belongs to schema edge, not sales"
    assert {"units", "notes"} <= set(yb._count_tables(config, "bronze_postgres"))


# --- create-if-missing: a run without --yaml_generation creates a named table's YAML, and leaves the YAML that exists alone

def test_create_only_writes_a_missing_yaml_and_leaves_an_existing_one_as_it_is(tmp_path):
    config = config_dir_for(tmp_path, "historical")
    run_generate(make_args("historical", tables=("notes",)), config)
    before = (config / "data_validation" / "notes.yaml").read_text(encoding="utf-8")
    results, _ = run_generate(make_args("historical", tables=("notes", "units")), config, create_only=True)
    got = statuses(results)
    assert got[("notes", "data_validation")] == "PRESENT" and got[("notes", "count_validation")] == "PRESENT"
    assert got[("units", "data_validation")] == "NEW" and got[("units", "count_validation")] == "NEW"
    assert (config / "data_validation" / "notes.yaml").read_text(encoding="utf-8") == before


def test_create_only_writes_the_count_entry_when_only_that_is_missing(tmp_path):
    config = config_dir_for(tmp_path, "historical")
    run_generate(make_args("historical", tables=("notes",), count="no"), config)
    results, _ = run_generate(make_args("historical", tables=("notes",)), config, create_only=True)
    got = statuses(results)
    assert got[("notes", "data_validation")] == "PRESENT"
    assert got[("notes", "count_validation")] == "NEW"


def test_tables_missing_yaml_lists_only_the_named_tables_with_a_missing_file_or_entry(tmp_path, monkeypatch):
    monkeypatch.setattr(generate_yaml, "FRAMEWORK_ROOT", tmp_path)
    config = config_dir_for(tmp_path, "historical")
    run_generate(make_args("historical", tables=("notes",)), config)
    assert generate_yaml.tables_missing_yaml("bronze_postgres", "historical", ["notes", "units"], "yes", "yes") == ["units"]
    assert generate_yaml.tables_missing_yaml("bronze_postgres", "historical", ["notes", "units"], "no", "no") == []


def test_a_missing_source_table_is_not_created_in_create_only_mode(tmp_path):
    config = config_dir_for(tmp_path, "historical")
    gone = FakeDB({k: v for k, v in SOURCE_TABLES.items() if k[1] != "units"}, PRIMARY_KEYS)
    results, _ = run_generate(make_args("historical", tables=("units",)), config, source=gone, create_only=True)
    assert statuses(results)[("units", "table")] == "FAILED"
    assert not (config / "data_validation" / "units.yaml").exists()


# --- incremental runs: the window, and a table without the column ----------------------------------------------

def test_an_incremental_run_writes_the_window_on_both_sides_of_data_and_count(tmp_path):
    config = config_dir_for(tmp_path, "incremental")
    results, _ = run_generate(make_args("incremental"), config)
    assert statuses(results)[("units", "data_validation")] == "NEW"
    block = data_block(config, "units")
    assert '"updated_at" > :start_time AND "updated_at" <= :end_time' in block["sourcequery"]
    assert '"UPDATED_AT" > :start_time AND "UPDATED_AT" <= :end_time' in block["targetquery"]
    count = count_block(config, "bronze_postgres", "units")
    assert '"updated_at" > :start_time AND "updated_at" <= :end_time' in count["sourcequery"]
    assert '"UPDATED_AT" > :start_time AND "UPDATED_AT" <= :end_time' in count["targetquery"]


def test_a_historical_run_has_no_window(tmp_path):
    config = config_dir_for(tmp_path, "historical")
    run_generate(make_args("historical"), config)
    assert ":start_time" not in data_block(config, "units")["sourcequery"]
    assert ":start_time" not in count_block(config, "bronze_postgres", "units")["sourcequery"]


def test_an_incremental_table_without_the_column_is_failed_for_data_and_count(tmp_path):
    config = config_dir_for(tmp_path, "incremental")
    results, _ = run_generate(make_args("incremental", tables=("lookup",)), config)
    assert statuses(results) == {("lookup", "data_validation"): "FAILED", ("lookup", "count_validation"): "FAILED"}
    assert detail_of(results, "lookup", "data_validation").startswith(
        "incremental column updated_at is not compared on both sides")
    assert not (config / "data_validation" / "lookup.yaml").exists()


def test_the_rules_workbook_can_name_the_incremental_column(tmp_path):
    config = config_dir_for(tmp_path, "incremental")
    run_generate(make_args("incremental"), config, table_rules={"units": TableRule(incremental_column="created_at")})
    assert '"created_at" > :start_time' in data_block(config, "units")["sourcequery"]
    assert '"CREATED_AT" > :start_time' in data_block(config, "units")["targetquery"]


def test_run_for_validation_passes_the_run_type_so_the_incremental_yaml_has_the_window(run_env):
    _run(run_env, run_type="incremental", tables=["units"])
    block = data_block(config_dir_for(run_env.root, "incremental"), "units")
    assert ":start_time" in block["sourcequery"] and ":end_time" in block["targetquery"]


DOCS_SOURCE = {("edge", "docs"): [("id", "integer"), ("meta", "jsonb"), ("tags", "hstore"),
                                  ("updated_at", "timestamp without time zone")]}
DOCS_TARGET = {("EDGE", "DOCS"): [("ID", "NUMBER"), ("META", "VARIANT"), ("TAGS", "TEXT"), ("UPDATED_AT", "TIMESTAMP_NTZ")]}


def test_jsonb_and_hstore_columns_are_listed_for_the_validator_and_selected_as_text(tmp_path):
    config = config_dir_for(tmp_path, "historical")
    source, target = FakeDB(DOCS_SOURCE, {("edge", "docs"): ["id"]}), FakeDB(DOCS_TARGET)
    results, _ = run_generate(make_args("historical", tables=("docs",)), config, source=source, target=target)
    assert statuses(results)[("docs", "data_validation")] == "NEW"
    block = data_block(config, "docs")
    assert block["json_columns"] == {"meta_normalized": "jsonb", "tags_normalized": "hstore"}
    assert 'CAST("meta" AS TEXT)' in block["sourcequery"]
    assert 'CAST("META" AS STRING)' in block["targetquery"]
    assert "# Compared as flattened JSON text (json_columns, both sides): meta_normalized, tags_normalized" in \
        (config / "data_validation" / "docs.yaml").read_text(encoding="utf-8")
