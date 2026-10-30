"""Row filters read from the Excel files in yaml_generation/Filter-files-excel/.

Why: a filter keeps part of a large table, so the framework does not load the whole table into RAM. The filter is
applied to the source and to Snowflake, for the count and the data, the same way as a filter in the rules workbook
(settings.py). When the rules workbook sets a filter for the same side, that one is used.

What is read
  A workbook with the columns "Target Table Name" and the filter column for the run type (the PostgreSQL layout):
    --run_type incremental  ->  "Incremental filters to apply"
    --run_type historical   ->  "Historical Filter to apply."
  Headers are matched trimmed and in lower case. It is read from the first sheet that has both columns.
  Each cell holds the table, its alias, the joins and the conditions, one clause per line:

      general_ledger_line_items j
      join facilities f on j.facility_id = f.id
      where f.company_id in (137,9088,501) and j.created_at >= '2024-08-01'

  The cell does not start with "select count(*) from": the framework adds it. A cell that does start with "select"
  is read as it is.

  It becomes one condition per side. Each join becomes a sub-select on the foreign key:

      created_at >= '2024-08-01' AND facility_id IN (SELECT id FROM "public"."facilities" WHERE company_id IN (137,9088,501))

  The Snowflake side is built from the same cell, so both sides keep the same rows.

Shapes read
  - select count(*) from <table> [alias], then inner joins, then an optional where clause joined with AND.
  - Each join is  <table> [alias] on <a>.<column> = <alias>.id, where <a> is already in the statement.
  - A condition uses one table. A bare column belongs to the first table. Joins chain (a table joined to a joined table).

What is refused, and the table is then loaded in full (the log names the row)
  OR, BETWEEN, left/right/outer joins, subqueries, group by, quoted identifiers, unbalanced quotes or brackets,
  a condition that uses two tables, and anything check_filter() refuses.

Other workbooks in the folder, for example Sitelink_bronze_filters.xlsx, have no SQL. They are named in the log and
not read. The statements are PostgreSQL, so they are used for PostgreSQL runs only.

Self-check: python yaml_generation/filter_files.py
"""

import logging
import re
import sys
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import load_workbook

FRAMEWORK_ROOT = Path(__file__).resolve().parents[1]
if str(FRAMEWORK_ROOT) not in sys.path:  # lets the self-check run as a plain script
    sys.path.insert(0, str(FRAMEWORK_ROOT))

from yaml_generation.db_rules import RULES, quote  # noqa: E402
from yaml_generation.yaml_builder import FIVETRAN_FILTER, check_filter  # noqa: E402

FOLDER = Path(__file__).resolve().parent / "Filter-files-excel"
# Header text, trimmed and in lower case. The filter column is the one for the run type.
TABLE_COLUMN = "target table name"
FILTER_COLUMNS = {"incremental": "incremental filters to apply", "historical": "historical filter to apply."}

# A quoted literal. Text inside quotes is never read as SQL, and never changed.
LITERAL = re.compile(r"'[^']*'")
# alias.column, as written in a condition.
QUALIFIED = re.compile(r"\b([A-Za-z_]\w*)\.(\w+)\b")
# The table in a FROM or JOIN part, with an optional alias.
TABLE = re.compile(r"^(\w+)(?:\s+(\w+))?$")
# The table a count statement is about: the name right after 'from'.
BASE_TABLE = re.compile(r"select\s+count\s*\(\s*\*\s*\)\s+from\s+(\w+)", re.IGNORECASE)
# One join: <table> [alias] on <a>.<column> = <b>.<column>.
JOIN_ON = re.compile(r"^(\w+)(?:\s+(\w+))?\s+on\s+(\w+)\.(\w+)\s*=\s*(\w+)\.(\w+)$", re.IGNORECASE)
# Words that change what the count means. Checked outside quotes only.
REFUSED_WORDS = re.compile(
    r"\b(or|between|exists|left|right|full|outer|cross|group|order|having|limit|union|distinct|select)\b",
    re.IGNORECASE)


class Refused(ValueError):
    """A statement this reader does not understand. The message says which part."""


@dataclass
class _Plan:
    """The tables and conditions of one statement, before they are written for a side."""

    base: str                                          # alias of the first table
    tables: dict = field(default_factory=dict)         # alias -> table name, as written
    own: dict = field(default_factory=dict)            # alias -> conditions that use only that table
    joins: dict = field(default_factory=dict)          # alias -> [(column, parent alias)], one sub-select each


@dataclass
class _Side:
    """How one side writes its tables. active is added inside every sub-select on that side."""

    ref: object                  # function: table name -> the name written after FROM / IN (SELECT ... FROM ...)
    active: str = None


def _outside_quotes(text, change):
    """Apply change() to the parts of text outside quotes. Quoted text is kept as it is."""
    parts = re.split(r"('[^']*')", text)
    return "".join(part if i % 2 else change(part) for i, part in enumerate(parts))


def _masked(text):
    """The text with each quote's contents removed, so a keyword or a dot inside quotes is not read."""
    return LITERAL.sub("''", text)


def _split_and(text):
    """Split a WHERE clause at each AND that is outside quotes."""
    pieces, start = [], 0
    for match in re.finditer(r"'[^']*'|\band\b", text, re.IGNORECASE):
        if match.group(0).startswith("'"):
            continue
        pieces.append(text[start:match.start()])
        start = match.end()
    pieces.append(text[start:])
    return [piece.strip() for piece in pieces]


def _plan(statement) -> _Plan:
    """Read one count statement into a _Plan. Raises Refused with the reason."""
    # Whitespace is collapsed outside quotes only, so a literal such as 'a  b' keeps its spaces.
    text = _outside_quotes(statement.strip(), lambda s: re.sub(r"\s+", " ", s)).strip()
    if text.endswith(";"):
        text = text[:-1].strip()
    if text.count("'") % 2:
        raise Refused("unbalanced quote")
    match = re.match(r"select\s+count\s*\(\s*\*\s*\)\s+from\s+(.+)$", text, re.IGNORECASE)
    if not match:
        raise Refused("only 'select count(*) from ...' statements are read")
    body = match.group(1)
    masked = _masked(body)
    if '"' in masked:
        raise Refused("quoted identifiers are not read")
    if masked.count("(") != masked.count(")"):
        raise Refused("brackets do not balance")
    word = REFUSED_WORDS.search(masked)
    if word:
        raise Refused(f"'{word.group(0)}' is not read")

    parts = re.split(r"\s+where\s+", body, maxsplit=1, flags=re.IGNORECASE)
    from_text, where_text = parts[0], (parts[1] if len(parts) > 1 else None)
    join_parts = re.split(r"\s+(?:inner\s+)?join\s+", from_text, flags=re.IGNORECASE)

    base = TABLE.match(join_parts[0].strip())
    if not base:
        raise Refused(f"the table is not read: '{join_parts[0].strip()}'")
    alias = (base.group(2) or base.group(1)).lower()
    plan = _Plan(base=alias)
    plan.tables[alias] = base.group(1)
    plan.own[alias] = []

    for piece in join_parts[1:]:
        join = JOIN_ON.match(piece.strip())
        if not join:
            raise Refused(f"join not read: '{piece.strip()}'. Write it as: join <table> <alias> on <alias>.<column> = <alias>.id")
        new_table, new_alias = join.group(1), (join.group(2) or join.group(1)).lower()
        a_alias, a_col, b_alias, b_col = (join.group(i).lower() for i in (3, 4, 5, 6))
        if new_alias in plan.tables:
            raise Refused(f"alias '{new_alias}' is used twice")
        # The joined table must be the one with the id. The other side is a table already in the statement.
        if a_alias == new_alias and a_col == "id" and b_alias in plan.tables:
            old, column = b_alias, b_col
        elif b_alias == new_alias and b_col == "id" and a_alias in plan.tables:
            old, column = a_alias, a_col
        else:
            raise Refused(f"join '{piece.strip()}' must link a table already in the statement to {new_table}.id")
        plan.tables[new_alias] = new_table
        plan.own[new_alias] = []
        plan.joins.setdefault(old, []).append((column, new_alias))

    for condition in (_split_and(where_text) if where_text else []):
        if not condition:
            raise Refused("an empty condition after WHERE or AND")
        masked = _masked(condition)
        if masked.count("(") != masked.count(")"):
            raise Refused(f"brackets do not balance in '{condition}'")
        used = {alias_name.lower() for alias_name, _ in QUALIFIED.findall(masked)}
        unknown = used - set(plan.tables)
        if unknown:
            raise Refused(f"'{sorted(unknown)[0]}' is not a table or alias in the statement: '{condition}'")
        if len(used) > 1:
            raise Refused(f"a condition uses two tables; write it as a join: '{condition}'")
        owner = used.pop() if used else plan.base
        # The alias is dropped: inside a sub-select the column belongs to that sub-select's table.
        plan.own[owner].append(_outside_quotes(
            condition, lambda s: re.sub(rf"\b{owner}\.(?=\w)", "", s, flags=re.IGNORECASE)))
    return plan


def _where(plan: _Plan, alias: str, side: _Side) -> list:
    """The conditions for one table: its own, then one sub-select for each join from it to a parent table."""
    parts = list(plan.own[alias])
    for column, parent in plan.joins.get(alias, []):
        inner = _where(plan, parent, side)
        if side.active:
            inner.append(side.active)
        sub = f"SELECT id FROM {side.ref(plan.tables[parent])}"
        if inner:
            sub += " WHERE " + " AND ".join(inner)
        parts.append(f"{column} IN ({sub})")
    return parts


def convert(statement, source_type, source_schema, target_database, target_schema):
    """(source_where, target_where) for one count statement. A side is None when the statement adds no condition.

    Raises Refused when the statement is not read. Tables without a schema are read in source_schema (source) and
    target_database.target_schema (Snowflake), as the rest of the generator does.
    """
    plan = _plan(statement)
    pg, snow = RULES[source_type], RULES["snowflake"]
    source = _Side(ref=lambda table: f"{quote(source_schema, pg)}.{quote(table, pg)}")
    # Snowflake keeps history rows, so a sub-select reads only the active ones (the rule the generator uses for the base table).
    # ponytail: no check that the join and condition columns exist. A wrong column fails at the count with a SQL error.
    target = _Side(ref=lambda table: f"{target_database}.{quote(target_schema, snow)}.{quote(table.upper(), snow)}",
                   active=FIVETRAN_FILTER)
    return _join(_where(plan, plan.base, source)), _join(_where(plan, plan.base, target))


def _join(parts):
    return " AND ".join(parts) if parts else None


def _statements(path, filter_column):
    """(row, table, statement) for each row of the first sheet that has both columns. None when no sheet has them."""
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        for sheet in workbook.worksheets:
            rows = sheet.iter_rows(values_only=True)
            header = [str(name).strip().lower() if name is not None else "" for name in (next(rows, None) or ())]
            if TABLE_COLUMN not in header or filter_column not in header:
                continue
            table_at, filter_at = header.index(TABLE_COLUMN), header.index(filter_column)
            found = []
            for row_number, row in enumerate(rows, start=2):
                table = row[table_at] if table_at < len(row) else None
                statement = row[filter_at] if filter_at < len(row) else None
                if table is not None and str(table).strip():
                    found.append((row_number, str(table).strip(), str(statement) if statement else ""))
            return found
        return None
    finally:
        workbook.close()


def _count_statement(text):
    """The cell as a count statement. The cell holds no 'select count(*) from' (the framework adds it), so the parser reads it as before."""
    text = text.strip()
    return text if text.lower().startswith("select") else "select count(*) from " + text


def load(source_type, source_schema, target_database, target_schema, run_type, folder=FOLDER, log=None) -> dict:
    """{table: (source_where, target_where)} for the tables whose cell was read. Tables not in it are loaded in full.

    run_type picks the column: incremental reads "Incremental filters to apply", historical reads "Historical Filter to apply.".
    target_database is the Snowflake name as written in the YAML (e.g. {env}_EDGE_BRONZE), so the target side keeps the placeholder.
    A refused cell is logged with its file and row, and that table is loaded in full.
    """
    log = log or logging.getLogger(__name__)
    if source_type != "postgres":
        log.info("Filter files: none used for %s, because the statements are PostgreSQL", source_type)
        return {}

    column = FILTER_COLUMNS[run_type]
    filters, seen, used_files = {}, set(), []
    for path in sorted(Path(folder).glob("*.xlsx")):
        if path.name.startswith("~$"):
            continue
        rows = _statements(path, column)
        if rows is None:
            log.warning("Filter file %s has no '%s' column, so it is not read", path.name, column)
            continue
        used_files.append(path.name)
        for row_number, table, statement in rows:
            if not statement.strip():
                continue
            label = f"{path.name} row {row_number} ({table}), column '{column}'"
            if table in seen:
                log.warning("Filter not applied: %s. The table is listed twice; the first row is used", label)
                continue
            try:
                # The statement must be about the row's own table. A cell copied from another row would filter the wrong table.
                statement_table = BASE_TABLE.match(_count_statement(statement))
                if not statement_table or statement_table.group(1).lower() != table.lower():
                    raise Refused(f"the statement is about '{statement_table.group(1) if statement_table else '?'}', not '{table}'")
                source_where, target_where = convert(_count_statement(statement), source_type, source_schema,
                                                     target_database, target_schema)
                check_filter(source_where, label)
                check_filter(target_where, label, placeholders=True)
            except ValueError as exc:
                log.warning("Filter not applied, the table is loaded in full: %s: %s", label, exc)
                continue
            seen.add(table)
            if source_where or target_where:
                filters[table] = (source_where, target_where)
    log.info("Filter files: %d table filter(s) read for run type %s, column '%s', from %s",
             len(filters), run_type, column, ", ".join(used_files) or "none")
    return filters


if __name__ == "__main__":
    # Self-check, in the style of yaml_builder.py. Runs with no database and no workbook.
    args = ("postgres", "public", "DEV_EDGE_BRONZE", "STOREDGE_FMS_PUBLIC")

    # One join: the joined table becomes a sub-select on the foreign key. Own conditions come first.
    src, tgt = convert(
        "select count(*) from  general_ledger_line_items  j\njoin  facilities f on j.facility_id=f.id \n"
        "where f.company_id in (137,9088,501) and j.created_at >=  '2024-08-01';", *args)
    assert src == ("created_at >= '2024-08-01' AND facility_id IN "
                   '(SELECT id FROM "public"."facilities" WHERE company_id in (137,9088,501))'), src
    assert tgt == ("created_at >= '2024-08-01' AND facility_id IN "
                   '(SELECT id FROM DEV_EDGE_BRONZE."STOREDGE_FMS_PUBLIC"."FACILITIES" '
                   "WHERE company_id in (137,9088,501) AND _FIVETRAN_ACTIVE = TRUE)"), tgt

    # Chained joins: phone_numbers -> tenants -> facilities. Each level is a sub-select inside the one above.
    src, _ = convert(
        "select count(*) from phone_numbers  j join  tenants  t on j.owner_id=t.id join  facilities  f on t.facility_id=f.id "
        "where  j.owner_type in ( 'Tenant') and f.company_id in (137,9088,501)", *args)
    assert src == ("owner_type in ( 'Tenant') AND owner_id IN (SELECT id FROM \"public\".\"tenants\" WHERE "
                   'facility_id IN (SELECT id FROM "public"."facilities" WHERE company_id in (137,9088,501)))'), src

    # A plain condition on one table, and no filter at all.
    assert convert("select count(*) from account_balance_items   \nwhere   created_at >=  '2026-08-01';", *args)[0] \
        == "created_at >= '2026-08-01'"
    assert convert("select count(*) from invoiceable_items", *args) == (None, None)

    # A join with no condition on the joined table keeps the rows that have a parent (inner join).
    src, _ = convert("select count(*) from discount_plan_ledger_instances_line_items j join discount_plan_ledger_instances ss "
                     "on j.discount_plan_ledger_instance_id=ss.id where created_at >=  '2026-08-01';", *args)
    assert src == ("created_at >= '2026-08-01' AND discount_plan_ledger_instance_id IN "
                   '(SELECT id FROM "public"."discount_plan_ledger_instances")'), src

    # Refused: each one names the reason. The table is then loaded in full.
    for bad in [
        "select count(*) from transactions  j join facilities f on j.facility_id=f.id where f.company_id in (137) "
        "and j.created_at >=  2025-08-01';",                                           # unbalanced quote (a typo in the workbook)
        "select count(*) from events j join facilities f on j.facility_id=f.id where f.company_id = 1 or j.created_at > '2025-01-01'",
        "select count(*) from events j left join facilities f on j.facility_id=f.id where f.company_id = 1",
        "select count(*) from events j join facilities f on j.facility_id=f.id where j.created_at > f.created_at",
        "select count(*) from events j join facilities f on j.facility_id=f.name where f.company_id = 1",
        "delete from events where created_at > '2025-01-01'",
    ]:
        try:
            convert(bad, *args)
        except ValueError:
            continue
        raise AssertionError(f"must be refused: {bad}")

    # A comment marker is read through, and load() refuses it with check_filter(), as for any rules-workbook filter.
    try:
        check_filter(convert("select count(*) from events j where j.created_at > '2025-01-01' -- note", *args)[0], "x")
    except ValueError:
        pass
    else:
        raise AssertionError("-- must be refused by check_filter")

    # Literal text is not read as SQL: an 'or' or a dot inside quotes is kept.
    assert convert("select count(*) from events where status = 'x or y.z'", *args)[0] == "status = 'x or y.z'"

    # A workbook cell has no 'select count(*) from'; the framework adds it. A cell that starts with 'select' is kept.
    assert _count_statement("events j\njoin facilities f on j.facility_id=f.id\nwhere f.company_id in (137)") \
        == "select count(*) from events j\njoin facilities f on j.facility_id=f.id\nwhere f.company_id in (137)"
    assert _count_statement("select count(*) from events") == "select count(*) from events"
    print("filter_files self-check passed")
